import csv
import io

import pytest
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.app.schemas import CommitmentNumberItem, RecordResponse


def to_response(record, record_id):
    return RecordResponse.model_validate(
        {
            **record.model_dump(),
            "id": record_id,
            "document_id": "doc-a",
            "created_at": "2026-01-01T00:00:00Z",
            "updated_at": "2026-01-01T00:00:00Z",
        }
    )
from backend.app.services.export_service import records_to_csv
from backend.app.services.record_service import RecordService
from backend.tests.test_record_service import complete_record


@pytest.fixture
def client():
    yield TestClient(app)


@pytest.fixture
def seed(db_sessionmaker):
    def make(records, document_id="doc-a"):
        with db_sessionmaker() as session:
            return RecordService(session).save_records(records, document_id)

    return make


def test_row_per_commitment_repeats_patient_details():
    record = complete_record(
        page_number=1,
        commitment_numbers=[
            CommitmentNumberItem(value="0070", confidence="high", label_type="commitment"),
            CommitmentNumberItem(value="0081", confidence="high", label_type="order"),
        ],
    )
    record = to_response(record, "rec-1")

    rows = list(csv.reader(io.StringIO(records_to_csv([record]).decode("utf-8-sig"))))

    assert rows[0][0] == "מספר התחייבות"
    assert [row[0] for row in rows[1:]] == ["0070", "0081"]
    assert [row[1] for row in rows[1:]] == ["טופס 17", "מספר דרישה"]
    # Patient/sample details repeat identically on both rows.
    for row in rows[1:]:
        assert row[4] == "000000018"  # id keeps leading zeros
        assert row[5] == "000123"     # central number
        assert row[6] == "2026-03-01" # received date
        assert row[7] == "סורוקה"       # sender
        assert row[8] == "1"          # page number


def test_record_without_commitment_still_exports_one_row():
    record = complete_record(commitment_numbers=[])
    record = to_response(record, "rec-2")

    rows = list(csv.reader(io.StringIO(records_to_csv([record]).decode("utf-8-sig"))))

    assert len(rows) == 2
    assert rows[1][0] == "" and rows[1][1] == ""
    assert rows[1][4] == "000000018"


def test_export_endpoint_returns_csv_attachment(client, seed):
    seed([complete_record(page_number=1)])

    response = client.get("/api/records/export", params={"format": "csv"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[1][0] == "0070"


def test_export_filters_by_status(client, seed):
    record = complete_record(page_number=2)
    seed([complete_record(page_number=1), record])

    response = client.get("/api/records/export", params={"status": "approved"})

    assert response.status_code == 200
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert len(rows) == 1  # header only — nothing approved yet


def test_export_json_returns_row_objects(client, seed):
    seed([complete_record(page_number=1)])

    response = client.get("/api/records/export", params={"format": "json"})

    assert response.status_code == 200
    assert "autocommit-export.json" in response.headers["content-disposition"]
    rows = response.json()
    assert rows[0]["מספר התחייבות"] == "0070"
    assert rows[0]["מספר זהות"] == "000000018"


def test_export_xlsx_returns_real_workbook(client, seed):
    from openpyxl import load_workbook

    seed([complete_record(page_number=1)])

    response = client.get("/api/records/export", params={"format": "xlsx"})

    assert response.status_code == 200
    assert "autocommit-export.xlsx" in response.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.cell(row=1, column=1).value == "מספר התחייבות"
    assert sheet.cell(row=2, column=1).value == "0070"


def test_export_rejects_unknown_format(client):
    assert client.get("/api/records/export", params={"format": "pdf"}).status_code == 422
